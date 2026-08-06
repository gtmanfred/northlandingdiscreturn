import io
import uuid
from datetime import date as _date

import openpyxl

from app.repositories.disc import DiscRepository
from app.services.disc_id_backfill import BackfillReport, backfill_ids, blocking_reason


def _sheet(data_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Current"
    ws.append(["North Landing Discs Database"])
    ws.append(["Sorted by ..."])
    ws.append(["Name", "Phone", "Mfr", "Model", "Color", "Other",
               "Code", "Date found", "Date retuned", "Date contacted", "ID"])
    for r in data_rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ids(out_bytes):
    ws = openpyxl.load_workbook(io.BytesIO(out_bytes)).active
    return {r[0].row: r[0].value for r in ws.iter_rows(min_row=4, min_col=11, max_col=11)}


async def test_matched_row_receives_the_databases_id(db):
    disc = await DiscRepository(db).create(
        manufacturer="Innova", name="Teebird", colors=["white"],
        input_date=_date(2026, 6, 1),
    )
    data = _sheet([["?", None, "Innova", "Teebird", "white", None, None,
                    _date(2026, 6, 1), None, None, None]])
    out, report = await backfill_ids(data, db)
    assert _ids(out)[4] == str(disc.id)
    assert report.matched == 1
    assert report.generated == 0


async def test_unmatched_row_receives_a_fresh_uuid(db):
    data = _sheet([["?", None, "Innova", "Nothing", "pink", None, None,
                    _date(2026, 6, 1), None, None, None]])
    out, report = await backfill_ids(data, db)
    value = _ids(out)[4]
    assert uuid.UUID(value)
    assert report.generated == 1
    assert report.matched == 0


async def test_populated_id_is_never_modified(db):
    keep = str(uuid.uuid4())
    data = _sheet([["?", None, "Innova", "Teebird", "white", None, None,
                    _date(2026, 6, 1), None, None, keep]])
    out, report = await backfill_ids(data, db)
    assert _ids(out)[4] == keep
    assert report.already_populated == 1
    assert report.matched == 0
    assert report.generated == 0


async def test_dry_run_returns_no_bytes_but_a_full_report(db):
    data = _sheet([["?", None, "Innova", "Nothing", "pink", None, None,
                    _date(2026, 6, 1), None, None, None]])
    out, report = await backfill_ids(data, db, dry_run=True)
    assert out is None
    assert report.generated == 1


async def test_second_row_resolving_to_a_claimed_disc_is_left_blank(db):
    await DiscRepository(db).create(
        manufacturer="Innova", name="Teebird", colors=["white"],
        input_date=_date(2026, 6, 1),
    )
    row = ["?", None, "Innova", "Teebird", "white", None, None,
           _date(2026, 6, 1), None, None, None]
    out, report = await backfill_ids(_sheet([list(row), list(row)]), db)
    ids = _ids(out)
    assert ids[4] is not None
    assert ids[5] is None
    assert report.matched == 1
    assert [s["row"] for s in report.skipped] == [5]
    assert "already claimed" in report.skipped[0]["reason"]


async def test_rows_with_id_errors_are_skipped_not_overwritten(db):
    data = _sheet([["?", None, "Innova", "Teebird", "white", None, None,
                    _date(2026, 6, 1), None, None, "not-a-uuid"]])
    out, report = await backfill_ids(data, db)
    assert _ids(out)[4] == "not-a-uuid"
    assert [s["reason"] for s in report.skipped] == ["invalid ID"]


async def test_row_missing_a_date_still_gets_a_fresh_uuid(db):
    data = _sheet([["?", None, "Innova", "Teebird", "white", None, None,
                    None, None, None, None]])
    out, report = await backfill_ids(data, db)
    assert uuid.UUID(_ids(out)[4])
    assert report.generated == 1


def test_blocking_reason_blocks_when_generated_exceeds_matched():
    report = BackfillReport(matched=1, generated=2)
    reason = blocking_reason(report)
    assert reason is not None
    assert "1" in reason
    assert "2" in reason


def test_blocking_reason_does_not_block_when_matched_meets_or_exceeds_generated():
    assert blocking_reason(BackfillReport(matched=5, generated=5)) is None
    assert blocking_reason(BackfillReport(matched=5, generated=2)) is None


def test_blocking_reason_does_not_block_an_all_zero_report():
    assert blocking_reason(BackfillReport(matched=0, generated=0)) is None
